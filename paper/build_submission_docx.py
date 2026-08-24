#!/usr/bin/env python3
"""Build the editable Preprints.org submission manuscript from canonical Markdown."""

from __future__ import annotations

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
import reportlab
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


INK = "102A43"
MUTED = "52667A"
BLUE = "2563EB"
TEAL = "0F8B8D"
AMBER = "C76A00"
RED = "B42318"
PALE_BLUE = "EAF2FF"
PALE_TEAL = "E7F7F5"
PALE_AMBER = "FFF4E5"
PALE_RED = "FDECEC"
RULE = "D5DEE8"
TABLE_HEADER = "F4F6F9"

BUNDLED_FONT_ROOT = Path(reportlab.__file__).resolve().parent / "fonts"
MACOS_FONT_ROOT = Path("/System/Library/Fonts/Supplemental")


def figure_font(macos_name, bundled_name):
    macos_path = MACOS_FONT_ROOT / macos_name
    return macos_path if macos_path.exists() else BUNDLED_FONT_ROOT / bundled_name


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    if sum(widths_dxa) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA, received {sum(widths_dxa)}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bottom_border(paragraph, color=RULE, size="6", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_code_frame(style):
    """Apply a restrained shaded frame to the dedicated code paragraph style."""
    p_pr = style._element.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), TABLE_HEADER)

    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), RULE)
        p_bdr.append(border)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, INK, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    custom = {
        "Manuscript Title": (20.5, INK, True, 0, 5, WD_ALIGN_PARAGRAPH.LEFT),
        "Author": (11.5, INK, True, 0, 2, WD_ALIGN_PARAGRAPH.LEFT),
        "Affiliation": (8.7, MUTED, False, 0, 1, WD_ALIGN_PARAGRAPH.LEFT),
        "Version Line": (8.5, MUTED, True, 3, 5, WD_ALIGN_PARAGRAPH.LEFT),
        "Abstract Heading": (13, BLUE, True, 6, 3, WD_ALIGN_PARAGRAPH.LEFT),
        "Abstract Text": (10, INK, False, 0, 3, WD_ALIGN_PARAGRAPH.LEFT),
        "Keywords": (8.7, MUTED, False, 3, 6, WD_ALIGN_PARAGRAPH.LEFT),
        "Caption": (8.5, MUTED, False, 4, 8, WD_ALIGN_PARAGRAPH.LEFT),
        "Reference": (8.5, INK, False, 0, 3, WD_ALIGN_PARAGRAPH.LEFT),
        "URL": (8.5, BLUE, False, 0, 6, WD_ALIGN_PARAGRAPH.LEFT),
        "Table Text": (8.5, INK, False, 0, 0, WD_ALIGN_PARAGRAPH.LEFT),
        "Table Header": (8.5, INK, True, 0, 0, WD_ALIGN_PARAGRAPH.LEFT),
    }
    for name, (size, color, bold, before, after, alignment) in custom.items():
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.line_spacing = 1.15 if name != "Reference" else 1.1
        if name == "Abstract Text":
            style.paragraph_format.line_spacing = 1.1
            style.paragraph_format.widow_control = True
        if name == "Reference":
            style.paragraph_format.left_indent = Inches(0.2)
            style.paragraph_format.first_line_indent = Inches(-0.2)
            style.paragraph_format.keep_together = True
            style.paragraph_format.widow_control = True

    code_style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "DejaVu Sans Mono"
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        code_style._element.rPr.rFonts.set(qn(f"w:{attr}"), "DejaVu Sans Mono")
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor.from_string(INK)
    code_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code_style.paragraph_format.left_indent = Inches(0.12)
    code_style.paragraph_format.right_indent = Inches(0.12)
    code_style.paragraph_format.space_before = Pt(6)
    code_style.paragraph_format.space_after = Pt(7)
    code_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    code_style.paragraph_format.line_spacing = Pt(11)
    # Fenced blocks are emitted as one paragraph with manual line breaks. Keep
    # the paragraph on one page when it fits so a lone first line is not left
    # at the bottom of the preceding page.
    code_style.paragraph_format.keep_together = True
    code_style.paragraph_format.widow_control = False
    add_code_frame(code_style)


def create_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "279")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.extend([fonts, color])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def create_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "right")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "600")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "600")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_bullet(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)


def add_hyperlink(paragraph, url, display_text, default_size):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = paragraph.add_run(display_text)
    set_run_font(run, size=default_size, color=BLUE)
    run.font.underline = True
    paragraph._p.remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, default_size=11):
    pattern = re.compile(
        r"(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|https?://\S+)"
    )
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=default_size, color=INK)
        token = match.group(0)
        if token.startswith("["):
            link_match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if link_match is None:
                raise ValueError(f"Malformed Markdown link: {token}")
            add_hyperlink(paragraph, link_match.group(2), link_match.group(1), default_size)
        elif token.startswith("http"):
            trailing = token[len(token.rstrip(".,;:")) :]
            url = token[: len(token) - len(trailing)] if trailing else token
            add_hyperlink(paragraph, url, url, default_size)
            if trailing:
                run = paragraph.add_run(trailing)
                set_run_font(run, size=default_size, color=INK)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="DejaVu Sans Mono", size=max(8.5, default_size - 1), color=INK)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=INK, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=default_size, color=INK, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=default_size, color=INK)


def add_box(draw, rect, fill, title, lines, font_bold, font_regular, border=RULE):
    x, y, w, h = rect
    draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=f"#{fill}", outline=f"#{border}", width=3)
    draw.text((x + 22, y + 18), title, font=font_bold, fill=f"#{INK}")
    for idx, line in enumerate(lines):
        draw.text((x + 22, y + 58 + idx * 30), line, font=font_regular, fill=f"#{MUTED}")


def draw_arrow(draw, start, end, color, dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        segments = 8
        for idx in range(segments):
            if idx % 2 == 0:
                xa = x1 + (x2 - x1) * idx / segments
                ya = y1 + (y2 - y1) * idx / segments
                xb = x1 + (x2 - x1) * (idx + 1) / segments
                yb = y1 + (y2 - y1) * (idx + 1) / segments
                draw.line((xa, ya, xb, yb), fill=f"#{color}", width=5)
    else:
        draw.line((x1, y1, x2, y2), fill=f"#{color}", width=5)
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 >= x1 else -1
        points = [(x2, y2), (x2 - sign * 22, y2 - 14), (x2 - sign * 22, y2 + 14)]
    else:
        sign = 1 if y2 >= y1 else -1
        points = [(x2, y2), (x2 - 14, y2 - sign * 22), (x2 + 14, y2 - sign * 22)]
    draw.polygon(points, fill=f"#{color}")


def build_figure(kind, output):
    output.parent.mkdir(parents=True, exist_ok=True)
    width, height = (1950, 760) if kind == "architecture" else (1950, 590)
    im = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(im)
    regular = ImageFont.truetype(str(figure_font("Arial.ttf", "Vera.ttf")), 25)
    small = ImageFont.truetype(str(figure_font("Arial.ttf", "Vera.ttf")), 22)
    bold = ImageFont.truetype(str(figure_font("Arial Bold.ttf", "VeraBd.ttf")), 27)
    italic = ImageFont.truetype(str(figure_font("Arial Italic.ttf", "VeraIt.ttf")), 22)
    label = ImageFont.truetype(str(figure_font("Arial Bold.ttf", "VeraBd.ttf")), 23)
    draw.rounded_rectangle((4, 4, width - 4, height - 4), radius=28, fill="#FBFCFE", outline=f"#{RULE}", width=4)

    if kind == "architecture":
        draw.text((55, 48), "ORDINARY DASHBOARD PIPELINE", font=label, fill=f"#{TEAL}")
        add_box(draw, (55, 135, 380, 185), PALE_BLUE, "Codex local state", ["candidate task rows", "each row's own rollout"], bold, small)
        add_box(draw, (535, 135, 380, 185), PALE_TEAL, "Ordinary reader", ["query-only SQL", "bounded rollout"], bold, small)
        add_box(draw, (1015, 135, 380, 185), PALE_BLUE, "Continuity policy", ["per-item evidence", "ranking suppression"], bold, small)
        add_box(draw, (1495, 135, 390, 185), PALE_TEAL, "Local interface", ["lifecycle controls", "codex:// deep link"], bold, small)
        draw_arrow(draw, (435, 228), (535, 228), TEAL)
        draw_arrow(draw, (915, 228), (1015, 228), TEAL)
        draw_arrow(draw, (1395, 228), (1495, 228), TEAL)
        draw.text((55, 372), "OPTIONAL WINGMAN PIPELINE", font=label, fill=f"#{AMBER}")
        add_box(draw, (55, 430, 380, 175), PALE_BLUE, "Codex local state", ["root + child rows", "rollout tails"], bold, small)
        add_box(draw, (535, 430, 380, 175), PALE_AMBER, "Wingman reader", ["task-tree graph", "bounded evidence"], bold, small)
        add_box(draw, (1015, 430, 380, 175), PALE_AMBER, "Preview + consent", ["bounded packet", "one-shot approval"], bold, small)
        add_box(draw, (1495, 430, 390, 175), PALE_RED, "External CLI", ["--ephemeral request", "service boundary"], bold, small, "F2B8B5")
        draw_arrow(draw, (435, 517), (535, 517), AMBER, dashed=True)
        draw_arrow(draw, (915, 517), (1015, 517), AMBER, dashed=True)
        draw_arrow(draw, (1395, 517), (1495, 517), RED, dashed=True)
        draw.text((55, 684), "The pipelines are separate. Requested read-only CLI mode is not OS isolation or a sole-context proof.", font=italic, fill=f"#{RED}")
    else:
        add_box(draw, (55, 120, 390, 190), PALE_BLUE, "Eligibility", ["not deferred", "complete history"], bold, small)
        add_box(draw, (535, 120, 390, 190), PALE_TEAL, "Fixed score", ["integer weights", "ID tie-break"], bold, small)
        add_box(draw, (1015, 120, 390, 190), PALE_BLUE, "Decision gate", ["top score >= 30", "lead >= 10 if runner-up"], bold, small)
        add_box(draw, (1495, 75, 390, 155), PALE_TEAL, "Recommend", ["bounded by input limit", "show reasons"], bold, small)
        add_box(draw, (1495, 290, 390, 155), PALE_AMBER, "Suppress ranking", ["no eligible item", "weak or close scores"], bold, small)
        draw_arrow(draw, (445, 215), (535, 215), TEAL)
        draw_arrow(draw, (925, 215), (1015, 215), TEAL)
        draw_arrow(draw, (1405, 190), (1495, 150), BLUE)
        draw_arrow(draw, (1405, 245), (1495, 365), AMBER)
        draw.text((1415, 125), "PASS", font=label, fill=f"#{BLUE}")
        draw.text((1410, 365), "FAIL", font=label, fill=f"#{AMBER}")
        draw.text((55, 515), "An incomplete item is excluded; it does not suppress a supported recommendation from another complete item.", font=italic, fill=f"#{MUTED}")
    im.save(output, dpi=(300, 300))


def setup_document():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)
    bullet_num_id = create_bullet_numbering(doc)
    decimal_num_id = create_decimal_numbering(doc)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("AIWINGMAN TECHNICAL NOTE")
    set_run_font(run, size=8, color=MUTED, bold=True)
    add_bottom_border(p, color=RULE, size="4", space="3")

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [7300, 2060])
    table.rows[0].cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)
    left_run = table.rows[0].cells[0].paragraphs[0].add_run("AiWingman Technical Note | Submission version 1.0 | 24 August 2026")
    set_run_font(left_run, size=8, color=MUTED)
    add_page_number(table.rows[0].cells[1].paragraphs[0])
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            borders.append(el)
        tc_pr.append(borders)
    return doc, bullet_num_id, decimal_num_id


def add_table(doc, raw_lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in raw_lines]
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    col_count = max(len(row) for row in rows)
    rows = [row + [""] * (col_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    if col_count == 4:
        if rows[0][0].startswith("Evidence"):
            widths = [2450, 1600, 1800, 3510]
        elif rows[0][0].startswith("Axis"):
            widths = [1400, 2750, 2500, 2710]
        else:
            widths = [2050, 2450, 2250, 2610]
    elif col_count == 3:
        if rows[0][0].startswith("Pipeline"):
            widths = [2050, 3000, 4310]
        else:
            widths = [2450, 2250, 4660]
    else:
        width = 9360 // col_count
        widths = [width] * (col_count - 1) + [9360 - width * (col_count - 1)]
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])

    for row_idx, row in enumerate(rows):
        prevent_row_split(table.rows[row_idx])
        for col_idx, value in enumerate(row):
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.style = "Table Header" if row_idx == 0 else "Table Text"
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, value, default_size=8.5)
            for run in paragraph.runs:
                if row_idx == 0:
                    run.bold = True
            if row_idx == 0:
                set_cell_shading(cell, TABLE_HEADER)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "FBFCFE")
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(0)


FENCE_OPEN_RE = re.compile(r"^(?P<marker>`{3,}|~{3,})(?:[ \t]*(?P<info>.*))?$")


def fenced_code_start(line):
    """Return the fence marker and optional info string for a Markdown code block."""
    match = FENCE_OPEN_RE.fullmatch(line.strip())
    if match is None:
        return None
    return match.group("marker"), (match.group("info") or "").strip()


def fenced_code_end(line, marker):
    """Recognize a closing fence of the same character and at least the same length."""
    candidate = line.strip()
    return (
        len(candidate) >= len(marker)
        and candidate[0] == marker[0]
        and candidate == candidate[0] * len(candidate)
    )


def read_fenced_code(lines, start):
    opening = fenced_code_start(lines[start])
    if opening is None:
        raise ValueError(f"Expected fenced code block at line {start + 1}")
    marker, info = opening
    code_lines = []
    i = start + 1
    while i < len(lines):
        if fenced_code_end(lines[i], marker):
            return code_lines, info, i + 1
        code_lines.append(lines[i].expandtabs(4))
        i += 1
    raise ValueError(f"Unclosed fenced code block at line {start + 1}")


def add_code_block(doc, code_lines):
    """Add one preformatted, left-aligned paragraph while preserving source line breaks."""
    paragraph = doc.add_paragraph(style="Code Block")
    run = paragraph.add_run()
    set_run_font(run, name="DejaVu Sans Mono", size=8.5, color=INK)
    if not code_lines:
        run.add_text(" ")
        return
    for index, code_line in enumerate(code_lines):
        if index:
            run.add_break(WD_BREAK.LINE)
        run.add_text(code_line)


def parse_manuscript(doc, markdown, bullet_num_id, decimal_num_id, figure_dir):
    lines = markdown.splitlines()
    nonempty = [idx for idx, line in enumerate(lines) if line.strip()]
    if len(nonempty) < 5 or not lines[nonempty[0]].startswith("# "):
        raise ValueError("Unexpected manuscript front matter")
    title_idx = nonempty[0]
    author_idx, affiliation_idx, orcid_idx, version_idx = nonempty[1:5]

    p = doc.add_paragraph(style="Manuscript Title")
    add_inline(p, lines[title_idx][2:].strip(), default_size=20.5)
    for style_name, idx, size in (
        ("Author", author_idx, 11.5),
        ("Affiliation", affiliation_idx, 8.7),
        ("Affiliation", orcid_idx, 8.7),
        ("Version Line", version_idx, 8.5),
    ):
        p = doc.add_paragraph(style=style_name)
        add_inline(p, lines[idx].strip(), default_size=size)

    i = version_idx + 1
    in_abstract = False
    decimal_numbering_used = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if fenced_code_start(lines[i]) is not None:
            code_lines, _language, i = read_fenced_code(lines, i)
            if doc.paragraphs and doc.paragraphs[-1].text.rstrip().endswith(":"):
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
            add_code_block(doc, code_lines)
            continue
        if line.startswith("<!-- FIGURE:"):
            kind = line[len("<!-- FIGURE:") :].split("-->")[0].strip()
            image_path = figure_dir / f"{kind}.png"
            build_figure(kind, image_path)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(str(image_path), width=Inches(6.45))
            doc_pr = run._r.xpath(".//wp:docPr")
            if doc_pr:
                doc_pr[0].set("descr", "AiWingman architecture diagram" if kind == "architecture" else "AiWingman conservative triage diagram")
            i += 1
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            style_name = "Abstract Heading" if heading == "Abstract" else "Heading 1"
            p = doc.add_paragraph(heading, style=style_name)
            in_abstract = heading == "Abstract"
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(line[4:].strip(), style="Heading 2")
            in_abstract = False
            i += 1
            continue
        if line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.208
                apply_bullet(p, bullet_num_id)
                add_inline(p, lines[i].strip()[2:].strip(), default_size=11)
                i += 1
            continue
        if re.match(r"^\d+\. ", line):
            # A Word numbering instance continues across non-adjacent lists when
            # the same numId is reused. Give every distinct Markdown list its
            # own instance so each list visibly restarts at 1.
            if decimal_numbering_used:
                list_num_id = create_decimal_numbering(doc)
            else:
                list_num_id = decimal_num_id
                decimal_numbering_used = True
            while i < len(lines) and re.match(r"^\d+\. ", lines[i].strip()):
                item = re.sub(r"^\d+\.\s+", "", lines[i].strip())
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.208
                apply_bullet(p, list_num_id)
                add_inline(p, item, default_size=11)
                i += 1
            continue

        parts = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith(("## ", "### ", "- ", "| ", "<!-- FIGURE:"))
                or re.match(r"^\d+\. ", nxt)
                or fenced_code_start(lines[i]) is not None
            ):
                break
            parts.append(nxt)
            i += 1
        text = " ".join(parts)
        if text.startswith(("Figure ", "Table ")):
            p = doc.add_paragraph(style="Caption")
            p.paragraph_format.keep_with_next = text.startswith("Figure ") is False
            add_inline(p, text, default_size=8.5)
        elif text.startswith("Keywords:"):
            p = doc.add_paragraph(style="Keywords")
            add_inline(p, text, default_size=8.7)
            in_abstract = False
        elif re.match(r"^\[\d+\] ", text):
            p = doc.add_paragraph(style="Reference")
            add_inline(p, text, default_size=8.5)
        elif re.fullmatch(r"https?://\S+", text):
            if doc.paragraphs and doc.paragraphs[-1].text.rstrip().endswith(":"):
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
            p = doc.add_paragraph(style="URL")
            add_inline(p, text, default_size=8.5)
        else:
            if in_abstract:
                p = doc.add_paragraph(style="Abstract Text")
                add_inline(p, text, default_size=10)
            else:
                p = doc.add_paragraph()
                add_inline(p, text, default_size=11)
                # Keep two semantically inseparable blocks intact in the final
                # submission rendering: the research-question lead-in with the
                # first numbered item, and the frozen-hash paragraph as a unit.
                if text.startswith("The study asks four questions."):
                    p.paragraph_format.keep_with_next = True
                if text.startswith("The specification SHA-256 is"):
                    p.paragraph_format.keep_together = True


def set_core_properties(doc):
    props = doc.core_properties
    props.title = (
        "AiWingman: A Local Continuity Overlay for Codex Task Portfolios with "
        "Retrospective Specification-Conformance Testing of Its Policy Layer"
    )
    props.subject = "AiWingman local continuity overlay and retrospective policy-layer conformance study"
    props.author = "Mehmet Solak"
    props.keywords = "coding agents; work continuity; local-first software; deterministic ranking suppression; specification-based testing; human oversight"
    props.comments = "AiWingman Technical Note - submission version 1.0"
    artifact_date = datetime(2026, 8, 24, tzinfo=timezone.utc)
    props.created = artifact_date
    props.modified = artifact_date


def remove_unused_custom_xml(doc):
    """Drop the empty template bibliography part and its generated package UUID."""
    for relationship_id, relationship in list(doc.part.rels.items()):
        if relationship.reltype.endswith("/customXml"):
            doc.part.drop_rel(relationship_id)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path(__file__).with_name("aiwingman_technical_report.md"))
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "output/docx/aiwingman-technical-note-v1.docx",
    )
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    figure_dir = Path(__file__).resolve().parents[1] / "tmp/docx-figures"
    doc, bullet_num_id, decimal_num_id = setup_document()
    set_core_properties(doc)
    parse_manuscript(
        doc,
        args.source.read_text(encoding="utf-8"),
        bullet_num_id,
        decimal_num_id,
        figure_dir,
    )
    remove_unused_custom_xml(doc)
    doc.save(args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
